"""Blocos de montagem do artigo em .docx (estilos, texto rico, equacoes, figuras, tabelas, citacoes).

O texto aceita **negrito**, *italico*, `codigo`, $matematica em LaTeX$ e citacoes [@chave] ou
[@chave1; @chave2]. As citacoes sao numeradas na ordem em que aparecem pela primeira vez.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from omml import omath, omath_para, set_math_size

CM_TEXT = 16.0
NAVY, BLUE, INK, INK_2 = "0D366B", "1C5CAB", "0B0B0B", "52514E"
BOX_COLORS = {
    "info": ("2A78D6", "EEF4FC"),
    "aviso": ("EDA100", "FDF6E3"),
    "chave": ("1BAF7A", "EAF7F1"),
}
INLINE_RE = re.compile(r"(\$[^$]+\$|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
CITE_RE = re.compile(r"\[@([^\]]+)\]")


class Article:
    def __init__(self, figures: list[str], tables: list[str], references: dict[str, str],
                 running_head: str) -> None:
        self.doc = Document()
        self.figures, self.tables, self.references = figures, tables, references
        self.figure_count = self.table_count = self.equation_count = 0
        self.equations: dict[str, int] = {}
        self.cited: dict[str, int] = {}
        self._setup(running_head)

    # ------------------------------------------------------------------ numeracao
    def fig(self, key: str) -> int:
        return self.figures.index(key) + 1

    def tab(self, key: str) -> int:
        return self.tables.index(key) + 1

    def eq(self, key: str) -> str:
        return f"({self.equations[key]})"

    # ------------------------------------------------------------------ estilos
    @staticmethod
    def _set_font(style, name: str) -> None:
        style.font.name = name
        properties = style.element.get_or_add_rPr()
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.insert(0, fonts)
        for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(attribute)) is not None:
                del fonts.attrib[qn(attribute)]
        for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attribute), name)

    def _style(self, name: str, base: str = "Normal"):
        try:
            return self.doc.styles[name]
        except KeyError:
            style = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = self.doc.styles[base]
            return style

    def _setup(self, running_head: str) -> None:
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin, section.bottom_margin = Cm(2.4), Cm(2.2)
        section.header_distance = section.footer_distance = Cm(1.1)

        normal = self.doc.styles["Normal"]
        self._set_font(normal, "Calibri")
        normal.font.size = Pt(11)
        normal.font.color.rgb = RGBColor.from_string(INK)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15

        for level, (size, color, before, after) in {
            1: (15, NAVY, 18, 6), 2: (12.5, BLUE, 12, 4), 3: (11.5, INK, 10, 3),
        }.items():
            style = self.doc.styles[f"Heading {level}"]
            self._set_font(style, "Calibri")
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.italic = False
            style.font.color.rgb = RGBColor.from_string(color)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.line_spacing = 1.0

        caption = self._style("Caption")
        self._set_font(caption, "Calibri")
        caption.font.size = Pt(9)
        caption.font.bold = False
        caption.font.italic = False
        caption.font.color.rgb = RGBColor.from_string(INK_2)
        caption.paragraph_format.space_before = Pt(2)
        caption.paragraph_format.space_after = Pt(10)
        caption.paragraph_format.line_spacing = 1.0

        for name in ("List Bullet", "List Number"):
            style = self.doc.styles[name]
            self._set_font(style, "Calibri")
            style.font.size = Pt(11)
            style.paragraph_format.space_after = Pt(3)
            style.paragraph_format.line_spacing = 1.15

        # A primeira pagina (bloco de titulo) fica sem cabecalho, como num periodico.
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run(running_head)
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(INK_2)
        for footer in (section.footer.paragraphs[0], section.first_page_footer.paragraphs[0]):
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._field(footer, "PAGE", "1", size=9)

    # ------------------------------------------------------------------ utilitarios
    @staticmethod
    def _field(paragraph, instruction: str, placeholder: str, size: float | None = None) -> None:
        def char(kind: str) -> None:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
            paragraph.add_run()._r.append(node)

        char("begin")
        node = OxmlElement("w:instrText")
        node.set(qn("xml:space"), "preserve")
        node.text = f" {instruction} "
        paragraph.add_run()._r.append(node)
        char("separate")
        shown = paragraph.add_run(placeholder)
        if size:
            shown.font.size = Pt(size)
        char("end")

    def _cite(self, text: str) -> str:
        def replace(match: re.Match) -> str:
            numbers = []
            for key in (part.strip().lstrip("@") for part in match.group(1).split(";")):
                if key not in self.references:
                    raise KeyError(f"referência desconhecida: {key}")
                self.cited.setdefault(key, len(self.cited) + 1)
                numbers.append(self.cited[key])
            return "[" + ", ".join(str(number) for number in sorted(numbers)) + "]"

        return CITE_RE.sub(replace, text)

    def rich(self, paragraph, text: str, size: float | None = None, color: str | None = None,
             bold: bool | None = None, italic: bool | None = None) -> None:
        for part in INLINE_RE.split(self._cite(text)):
            if not part:
                continue
            if part.startswith("$") and part.endswith("$") and len(part) > 1:
                element = omath(part[1:-1])
                if size:
                    set_math_size(element, size)
                paragraph._p.append(element)
                continue
            code = part.startswith("`") and part.endswith("`")
            if part.startswith("**") and part.endswith("**"):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            elif code:
                run = paragraph.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt((size or 11) - 1.5)
            else:
                run = paragraph.add_run(part)
            if size and not code:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
            if bold is not None and not part.startswith("**"):
                run.bold = bold
            if italic is not None:
                run.italic = italic

    @staticmethod
    def _insert_ppr(properties, *elements) -> None:
        # Ordem do esquema: pBdr e shd vem antes de tabs, spacing, ind e jc.
        for tag in ("w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc"):
            anchor = properties.find(qn(tag))
            if anchor is not None:
                for element in elements:
                    anchor.addprevious(element)
                return
        for element in elements:
            properties.append(element)

    def _shade(self, paragraph, border: str, fill: str) -> None:
        borders = OxmlElement("w:pBdr")
        for side, color, size in (("top", fill, "4"), ("left", border, "24"), ("bottom", fill, "4"),
                                  ("right", fill, "4")):
            node = OxmlElement(f"w:{side}")
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), size)
            node.set(qn("w:space"), "8" if side == "left" else "6")
            node.set(qn("w:color"), color)
            borders.append(node)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), fill)
        self._insert_ppr(paragraph._p.get_or_add_pPr(), borders, shading)

    def rule(self, color: str = "C3C2B7") -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(10)
        borders = OxmlElement("w:pBdr")
        node = OxmlElement("w:bottom")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "1")
        node.set(qn("w:color"), color)
        borders.append(node)
        self._insert_ppr(paragraph._p.get_or_add_pPr(), borders)

    # ------------------------------------------------------------------ texto
    def h1(self, text: str) -> None:
        self.doc.add_paragraph(text, style="Heading 1")

    def h2(self, text: str) -> None:
        self.doc.add_paragraph(text, style="Heading 2")

    def p(self, text: str, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size: float | None = None):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = align
        self.rich(paragraph, text, size=size)
        return paragraph

    def bullets(self, items: list[str], numbered: bool = False) -> None:
        for item in items:
            paragraph = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            self.rich(paragraph, item)

    def box(self, title: str, texts: list[str] | str, kind: str = "info") -> None:
        border, fill = BOX_COLORS[kind]
        texts = [texts] if isinstance(texts, str) else texts
        for index, text in enumerate(texts):
            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if index == 0 and title:
                run = paragraph.add_run(title + " ")
                run.bold = True
                run.font.size = Pt(10.5)
            self.rich(paragraph, text, size=10.5)
            fmt = paragraph.paragraph_format
            fmt.left_indent, fmt.right_indent = Cm(0.35), Cm(0.25)
            fmt.space_before = Pt(10 if index == 0 else 0)
            fmt.space_after = Pt(12 if index == len(texts) - 1 else 4)
            fmt.line_spacing = 1.1
            fmt.keep_together = True
            if index < len(texts) - 1:
                fmt.keep_with_next = True
            self._shade(paragraph, border, fill)

    def equation(self, key: str, latex: str) -> int:
        self.equation_count += 1
        self.equations[key] = self.equation_count
        table = self.doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        self._table_borders(table, None)
        widths = (CM_TEXT - 1.4, 1.4)
        for cell, width in zip(table.rows[0].cells, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            column.w = Cm(width)
        math_paragraph = table.rows[0].cells[0].paragraphs[0]
        math_paragraph.paragraph_format.space_before = Pt(3)
        math_paragraph.paragraph_format.space_after = Pt(3)
        math_paragraph._p.append(omath_para(latex))
        label = table.rows[0].cells[1].paragraphs[0]
        label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        label.paragraph_format.space_after = Pt(0)
        run = label.add_run(f"({self.equation_count})")
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(INK_2)
        return self.equation_count

    # ------------------------------------------------------------------ figuras e tabelas
    def figure(self, key: str, paths: list[Path] | Path, caption: str, width_cm: float = CM_TEXT) -> None:
        self.figure_count += 1
        expected = self.figures[self.figure_count - 1]
        if expected != key:
            raise RuntimeError(f"ordem das figuras: esperado {expected}, veio {key}")
        for path in [paths] if isinstance(paths, Path) else paths:
            if not path.is_file():
                raise FileNotFoundError(path)
            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
        paragraph = self.doc.add_paragraph(style="Caption")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self.rich(paragraph, f"**Figura {self.figure_count}** – {caption}", size=9)

    @staticmethod
    def _table_borders(table, specification) -> None:
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement(f"w:{side}")
            value = (specification or {}).get(side)
            if value is None:
                node.set(qn("w:val"), "nil")
            else:
                size, color = value
                node.set(qn("w:val"), "single")
                node.set(qn("w:sz"), size)
                node.set(qn("w:space"), "0")
                node.set(qn("w:color"), color)
            borders.append(node)
        properties = table._tbl.tblPr
        existing = properties.find(qn("w:tblBorders"))
        if existing is not None:
            properties.remove(existing)
        for tag in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption"):
            anchor = properties.find(qn(tag))
            if anchor is not None:
                anchor.addprevious(borders)
                return
        properties.append(borders)

    def table(self, key: str, caption: str, header: list[str], rows: list[list[str]],
              widths: list[float], align: str | None = None, size: float = 9.5,
              bold_rows: set[int] | None = None, note: str | None = None) -> None:
        self.table_count += 1
        expected = self.tables[self.table_count - 1]
        if expected != key:
            raise RuntimeError(f"ordem das tabelas: esperado {expected}, veio {key}")
        if sum(widths) > CM_TEXT + 0.05:
            raise ValueError(f"tabela {key} mais larga que o texto: {sum(widths)}")
        title = self.doc.add_paragraph(style="Caption")
        title.paragraph_format.space_before = Pt(8)
        title.paragraph_format.space_after = Pt(3)
        title.paragraph_format.keep_with_next = True
        self.rich(title, f"**Tabela {self.table_count}** – {caption}", size=9)

        align = align or "l" * len(header)
        table = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        self._table_borders(table, {"top": ("8", NAVY), "bottom": ("8", NAVY), "insideH": ("4", "D9D9D9")})
        for column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            column.w = Cm(width)
        alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
        bold_rows = bold_rows or set()
        for row_index, values in enumerate([header] + rows):
            row = table.rows[row_index]
            row_properties = row._tr.get_or_add_trPr()
            row_properties.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                repeat = OxmlElement("w:tblHeader")
                repeat.set(qn("w:val"), "true")
                row_properties.append(repeat)
            for column_index, (cell, value) in enumerate(zip(row.cells, values)):
                cell.width = Cm(widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                for line_index, line in enumerate(str(value).split("\n")):
                    if line_index:
                        paragraph = cell.add_paragraph()
                    paragraph.alignment = alignment[align[column_index]]
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.0
                    emphasis = row_index == 0 or (row_index - 1) in bold_rows
                    self.rich(paragraph, line, size=size, bold=True if emphasis else None)
                if row_index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:val"), "clear")
                    shading.set(qn("w:color"), "auto")
                    shading.set(qn("w:fill"), "E3EDF9")
                    # Ordem do esquema em tcPr: shd vem antes de vAlign.
                    cell_properties = cell._tc.get_or_add_tcPr()
                    v_align = cell_properties.find(qn("w:vAlign"))
                    if v_align is not None:
                        v_align.addprevious(shading)
                    else:
                        cell_properties.append(shading)
        if len(rows) <= 15:
            # Tabelas curtas ficam inteiras numa pagina.
            for row in table.rows[:-1]:
                for cell in row.cells:
                    for cell_paragraph in cell.paragraphs:
                        cell_paragraph.paragraph_format.keep_with_next = True
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(10 if note else 4)
        if note:
            self.rich(paragraph, note, size=8.5, color=INK_2)
        else:
            paragraph.paragraph_format.line_spacing = 0.8

    # ------------------------------------------------------------------ referencias e saida
    def reference_list(self) -> list[str]:
        unused = [key for key in self.references if key not in self.cited]
        for key, number in sorted(self.cited.items(), key=lambda item: item[1]):
            paragraph = self.doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Cm(0.8)
            paragraph.paragraph_format.first_line_indent = Cm(-0.8)
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.paragraph_format.line_spacing = 1.05
            self.rich(paragraph, f"[{number}] {self.references[key]}", size=9.5)
        return unused

    def check(self) -> None:
        if self.figure_count != len(self.figures) or self.table_count != len(self.tables):
            raise RuntimeError(f"figuras {self.figure_count}/{len(self.figures)}, "
                               f"tabelas {self.table_count}/{len(self.tables)}")

    def _modern_compatibility(self) -> None:
        # Sem isto o Word abre o arquivo em "Modo de Compatibilidade" (Word 2007).
        settings = self.doc.settings.element
        compat = settings.find(qn("w:compat"))
        if compat is None:
            compat = OxmlElement("w:compat")
            settings.append(compat)
        for node in compat.findall(qn("w:compatSetting")):
            if node.get(qn("w:name")) == "compatibilityMode":
                compat.remove(node)
        node = OxmlElement("w:compatSetting")
        node.set(qn("w:name"), "compatibilityMode")
        node.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
        node.set(qn("w:val"), "15")
        compat.append(node)

    def save(self, path: Path, title: str, subject: str, keywords: str, author: str) -> None:
        properties = self.doc.core_properties
        properties.title, properties.subject = title, subject
        properties.keywords, properties.author = keywords, author
        properties.language = "pt-BR"
        self._modern_compatibility()
        self.doc.save(path)
